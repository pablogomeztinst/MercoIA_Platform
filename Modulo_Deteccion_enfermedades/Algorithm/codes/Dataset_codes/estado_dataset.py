"""Genera estado_dataset.xlsx (raíz del repo) con el estado del dataset por fruta, contado desde disco:
  Resumen: fotos crudas -> recortes segmentados -> etiquetados -> pendientes -> en entrenamiento (+ gráfico)
  Mora / Pina / Banano: etiquetados por clase (etiquetas.csv) con torta, dataset de entrenamiento
  (Algorithm/data) con torta, y pendientes de etiquetar por sesión.

Uso:  python estado_dataset.py
"""
import csv
import re
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Font, PatternFill

AQUI = Path(__file__).resolve().parent
RAIZ = AQUI.parents[2]
CRUDAS, SEGM, DATA = AQUI.parents[1] / "Our_dataset_sin_segmen", RAIZ / "dataset_segm", AQUI.parents[1] / "data"
ETIQUETAS = AQUI / "data" / "etiquetas.csv"
IMG = {".png", ".jpg", ".jpeg"}
# hoja -> (carpeta crudas, carpeta dataset_segm, prefijo clases etiquetas, prefijo clases entrenamiento)
FRUTAS = {"Mora": ("Moras", "Moras", "mora_", "mora_"), "Pina": ("Pinas", "Pinas", "pina_", "pina_"),
          "Banano": ("bananos", "Bananos", "banano_", "banano_")}
NOMBRE = {"sana": "Sana", "sano": "Sano", "racimo_sano": "Sano", "antracnosis": "Antracnosis", "moho_gris": "Moho gris",
          "botritis": "Botritis", "enferma": "Enferma", "dano_externo": "Daño externo", "corte_interno": "Corte interno", "descartado": "Descartado", "fusarium": "Fusarium", "fusariosis": "Fusariosis", "pudricion": "Pudrición",
          "phytophthora": "Phytophthora", "phytopthora": "Phytophthora", "afectacion_externa": "Afectación externa"}
AZUL, GRIS = PatternFill("solid", fgColor="1F4E79"), PatternFill("solid", fgColor="F2F2F2")


def bonito(clase, prefijo):
  return NOMBRE.get(clase[len(prefijo):].lower(), clase)


def encabezado(ws, fila, textos):
  for j, t in enumerate(textos, 1):
    c = ws.cell(fila, j, t); c.font = Font(bold=True, color="FFFFFF"); c.fill = AZUL
    c.alignment = Alignment(horizontal="center", wrap_text=True)


def tabla(ws, fila, titulo, cabeceras, filas, fila_total=True):
  """Escribe título + tabla desde `fila`; devuelve (primera fila de datos, última fila de datos, fila total)."""
  ws.cell(fila, 1, titulo).font = Font(bold=True, size=12); fila += 1
  encabezado(ws, fila, cabeceras); ini = fila + 1
  for k, f in enumerate(filas):
    for j, v in enumerate(f, 1):
      ws.cell(ini + k, j, v)
  fin = ini + len(filas) - 1
  if fila_total:
    ws.cell(fin + 1, 1, "Total").font = Font(bold=True)
    for j in range(2, len(cabeceras) + 1):
      col = ws.cell(ini, j).column_letter
      if isinstance(ws.cell(ini, j).value, (int, float)) or str(ws.cell(ini, j).value).startswith("="):
        ws.cell(fin + 1, j, f"=SUM({col}{ini}:{col}{fin})").font = Font(bold=True)
  return ini, fin, fin + 1


def torta(ws, titulo, col_etq, col_val, ini, fin, ancla):
  ch = PieChart(); ch.title = titulo; ch.height, ch.width = 7.5, 11
  ch.add_data(Reference(ws, min_col=col_val, min_row=ini, max_row=fin))
  ch.set_categories(Reference(ws, min_col=col_etq, min_row=ini, max_row=fin))
  ch.dataLabels = etiquetas_datos(showPercent=True); ch.dataLabels.position = "bestFit"
  ws.add_chart(ch, ancla)


def etiquetas_datos(**activas):
  """DataLabelList con todo apagado salvo lo pedido (openpyxl deja serie y categoría encendidas)."""
  d = DataLabelList()
  for k in ("showVal", "showPercent", "showCatName", "showSerName", "showLegendKey", "showBubbleSize"):
    setattr(d, k, activas.get(k, False))
  d.showLeaderLines = True
  return d


def main():
  etq = {r["objeto"]: r["clase"] for r in csv.DictReader(open(ETIQUETAS, encoding="utf-8"))}
  hoy = date.today().isoformat()
  wb = Workbook(); res = wb.active; res.title = "Resumen"
  resumen = []

  for hoja, (crud, seg, pe, pt) in FRUTAS.items():
    ws = wb.create_sheet(hoja)
    titulo = {"Pina": "Piña", "Banano": "Racimo de banano"}.get(hoja, hoja)
    ws.cell(1, 1, f"Estado del dataset - {titulo}").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"Conteo de archivos en disco el {hoy}.").font = Font(italic=True, color="666666")

    fotos = [p for p in (CRUDAS / crud).iterdir() if p.suffix.lower() in IMG]
    recortes = [p for p in (SEGM / seg).iterdir() if p.suffix.lower() in IMG]
    lab = Counter(etq[f"{seg}/{p.name}"] for p in recortes if f"{seg}/{p.name}" in etq)
    pend = [p for p in recortes if f"{seg}/{p.name}" not in etq]
    ses = Counter(re.sub(r"_\d+(_\d{2,}|_mask\d+)?$", "", p.stem) for p in pend)
    entren = {d.name: ([f for f in d.iterdir() if not f.name.startswith("aug_")], [f for f in d.iterdir() if f.name.startswith("aug_")])
              for d in DATA.iterdir() if d.is_dir() and d.name.startswith(pt) and any(d.iterdir())}

    # Tabla 1: etiquetados por clase (fuente de verdad: etiquetas.csv)
    filas = [[bonito(c, pe), n, None] for c, n in lab.most_common()]
    ini, fin, tot = tabla(ws, 4, "1. Recortes etiquetados por clase (etiquetas.csv)", ["Clase", "Recortes", "%"], filas)
    for r in range(ini, fin + 1):
      ws.cell(r, 3, f"=B{r}/$B${tot}").number_format = "0.0%"
    torta(ws, f"{hoja}: etiquetados por clase", 1, 2, ini, fin, "G4")

    # Pendientes
    f = tot + 2
    ws.cell(f, 1, f"Recortes pendientes de etiquetar: {len(pend)}").font = Font(bold=True, size=12, color="C00000")
    top = ses.most_common(8)
    texto = "Por sesión: " + ", ".join(f"{s} ({n})" for s, n in top) + (f" y {len(ses) - 8} sesiones más" if len(ses) > 8 else "")
    ws.cell(f + 1, 1, texto).alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=f + 1, start_column=1, end_row=f + 1, end_column=5)
    ws.row_dimensions[f + 1].height = 15 * (len(texto) // 90 + 1)

    # Tabla 2: dataset de entrenamiento actual (Algorithm/data), como en la versión anterior del libro.
    # Empieza como mínimo en la fila 21 para que su torta no se monte sobre la de la tabla 1 (~15 filas).
    f = max(f + 4, 21)
    filas = [[bonito(c, pt), len(o), len(a), None, None] for c, (o, a) in sorted(entren.items(), key=lambda x: -len(x[1][0]))]
    ini2, fin2, tot2 = tabla(ws, f, "2. Dataset de entrenamiento actual (Algorithm/data)",
                             ["Clase", "Imágenes originales", "Imágenes aumentadas", "Total", "% originales"], filas)
    for r in range(ini2, fin2 + 1):
      ws.cell(r, 4, f"=B{r}+C{r}"); ws.cell(r, 5, f"=B{r}/$B${tot2}").number_format = "0.0%"
    torta(ws, f"{hoja}: entrenamiento (originales)", 1, 2, ini2, fin2, f"G{f}")
    ws.cell(tot2 + 2, 1, "Nota: la tabla 2 no se construyó desde etiquetas.csv (se armó a mano), por eso sus clases y "
            "conteos no coinciden con la tabla 1. Aumentadas = copias aug_* de generar_augmentacion.py.").alignment = Alignment(wrap_text=True)
    ws.merge_cells(start_row=tot2 + 2, start_column=1, end_row=tot2 + 2, end_column=5); ws.row_dimensions[tot2 + 2].height = 45
    for col, w in zip("ABCDE", (26, 18, 18, 10, 12)):
      ws.column_dimensions[col].width = w
    resumen.append([hoja, len(fotos), len(recortes), sum(lab.values()), len(pend), sum(len(o) for o, _ in entren.values())])

  # Hoja Resumen
  res.cell(1, 1, "Estado del dataset MERCOIA por fruta").font = Font(bold=True, size=14)
  res.cell(2, 1, f"Conteo de archivos en disco el {hoy}.").font = Font(italic=True, color="666666")
  ini, fin, tot = tabla(res, 4, "Embudo del dataset", ["Fruta", "Fotos crudas", "Recortes segmentados", "Recortes etiquetados",
                                                          "Pendientes de etiquetar", "En entrenamiento (originales)"], resumen)
  ch = BarChart(); ch.type = "col"; ch.grouping = "clustered"; ch.title = "Recortes por fruta"; ch.height, ch.width = 8, 16
  ch.add_data(Reference(res, min_col=4, max_col=5, min_row=ini - 1, max_row=fin), titles_from_data=True)
  ch.set_categories(Reference(res, min_col=1, min_row=ini, max_row=fin))
  ch.dataLabels = etiquetas_datos(showVal=True)
  res.add_chart(ch, "A11")
  res.cell(9, 1, "Fotos crudas = Our_dataset_sin_segmen. Recortes = dataset_segm (SAM 3). Etiquetados = filas de etiquetas.csv con archivo. "
           "Pendientes = recortes sin fila en etiquetas.csv. Entrenamiento = originales en Algorithm/data.").alignment = Alignment(wrap_text=True)
  res.merge_cells("A9:F9"); res.row_dimensions[9].height = 45
  for col, w in zip("ABCDEF", (12, 14, 20, 20, 22, 26)):
    res.column_dimensions[col].width = w

  salida = RAIZ / "estado_dataset.xlsx"
  wb.save(salida); print("guardado:", salida)
  for r in resumen:
    print(f"  {r[0]:7s} fotos {r[1]:4d} | recortes {r[2]:4d} | etiquetados {r[3]:4d} | PENDIENTES {r[4]:4d} | entrenamiento {r[5]:4d}")


if __name__ == "__main__":
  main()
