"""Parser del boletín "SIPSA Diario" (.xlsx) que publica el DANE.

Estructura real del archivo (verificada contra boletines reales descargados
de dane.gov.co/files/operaciones/SIPSA/anex-SIPSADiario-<dia><mes><año>.xlsx):

    Fila 1: vacía
    Fila 2: título con la fecha, ej. "Martes 11 de agosto de 2026"   -> índice 0
    Fila 3: nombre de cada mercado (columnas fusionadas), ej.
            [None, "Armenia, Mercar", None, "Bogotá, Corabastos", None, ...]
    Fila 4: subtítulos "Precio" / "Var %" repetidos por cada mercado
    Fila 5+: filas de categoría (ej. "Verduras y hortalizas", el resto vacío)
             intercaladas con filas de producto:
             ["Papa negra*", 1770, 0.0, 2900, -0.01, ...]

En vez de asumir posiciones fijas de columna (frágil ante cambios del DANE),
localizamos dinámicamente las columnas de "Bogotá, Corabastos" y
"Bucaramanga, Centroabastos" buscando esos textos en la fila de mercados.

El índice de cada fila (``index``) se cuenta a partir de la fila 2 (la del
título), igual que hacía n8n con su nodo "Extract From File": así el "x" que
devuelve el modelo normalizador sigue apuntando a la fila correcta.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field

import openpyxl

from app.services.text_utils import normalize_basic

EXCLUIDOS = [
    "martes", "precio kg", "precio", "verduras y hortalizas",
    "frutas frescas", "tuberculos raices y platanos",
    "variedad predominante", "variacion porcentual", "no disponible",
]


@dataclass
class FilaSipsa:
    index: int
    nombre: str | None
    precio_centroabastos: object = None
    var_centroabastos: object = None
    precio_corabastos: object = None
    var_corabastos: object = None
    # Solo se usa cuando la fila entera se "importó" de otra fuente/fecha (ver
    # sipsa_fetch._completar_con_fuente_principal): fecha real a la que
    # pertenecen precio_centroabastos/precio_corabastos, para que pricing.py
    # se lo aclare al usuario en vez de dar a entender que es del día del
    # informe. `None` en el caso normal (la fila sí es del día del boletín).
    nota_fecha: str | None = None


@dataclass
class BoletinSipsa:
    fecha_reporte_texto: str
    filas: list[FilaSipsa] = field(default_factory=list)


class SipsaParseError(Exception):
    pass


def _encontrar_columnas_mercado(ws) -> tuple[int, int, int | None, int | None]:
    """Devuelve (col_precio_corabastos, col_var_corabastos,
    col_precio_centroabastos, col_var_centroabastos), 0-indexadas.

    IMPORTANTE (descubierto comparando 5 boletines diarios reales): el DANE
    NO reporta siempre los mismos mercados regionales — rota entre
    Barranquilla, Bucaramanga, Cali, Cartagena, Medellín, Pasto, etc. según
    el día. Bogotá-Corabastos apareció en los 5 boletines revisados;
    Bucaramanga-Centroabastos solo en 2 de 5. Por eso Bogotá es obligatorio
    (si falta, el archivo realmente no tiene el formato esperado) y
    Bucaramanga es opcional: si ese día no viene, se devuelve `None` en vez
    de descartar el boletín completo (eso hacía que la búsqueda saltara
    innecesariamente más días atrás de lo debido para encontrar un boletín
    con ambos mercados). El respaldo del precio de Bucaramanga para ese caso
    se resuelve aparte, en sipsa_fetch.buscar_precio_respaldo().
    """
    col_corabastos = col_centroabastos = None

    # La fila de mercados suele ser la 3, pero se busca en una ventana amplia
    # por si el DANE llegara a insertar/quitar una fila.
    for row in ws.iter_rows(min_row=1, max_row=8):
        for cell in row:
            valor = normalize_basic(cell.value)
            if not valor:
                continue
            if "bogota" in valor and "corabastos" in valor:
                col_corabastos = cell.column - 1
            elif "bucaramanga" in valor and "centroabastos" in valor:
                col_centroabastos = cell.column - 1
        if col_corabastos is not None and col_centroabastos is not None:
            break

    if col_corabastos is None:
        raise SipsaParseError(
            "No se encontró la columna de Bogotá-Corabastos en el boletín SIPSA "
            "(ese mercado debería estar presente todos los días)."
        )

    # La columna de "Var %" siempre es la inmediatamente siguiente a "Precio".
    col_var_centroabastos = col_centroabastos + 1 if col_centroabastos is not None else None
    return col_corabastos, col_corabastos + 1, col_centroabastos, col_var_centroabastos


def parse_workbook(contenido: bytes) -> BoletinSipsa:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl lanza varias excepciones distintas
        raise SipsaParseError(f"El archivo descargado no es un .xlsx válido: {exc}") from exc

    ws = wb[wb.sheetnames[0]]
    col_pc, col_vc, col_pcen, col_vcen = _encontrar_columnas_mercado(ws)

    filas: list[FilaSipsa] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):

        def get(col: int | None):
            if col is None:
                return None
            return row[col] if col < len(row) else None

        filas.append(
            FilaSipsa(
                index=i,
                nombre=get(0),
                precio_corabastos=get(col_pc),
                var_corabastos=get(col_vc),
                precio_centroabastos=get(col_pcen),
                var_centroabastos=get(col_vcen),
            )
        )

    if not filas or not filas[0].nombre:
        raise SipsaParseError("El boletín SIPSA no tiene el formato esperado (falta título/fecha).")

    return BoletinSipsa(fecha_reporte_texto=str(filas[0].nombre), filas=filas)


def es_producto(fila: FilaSipsa) -> bool:
    nombre = normalize_basic(fila.nombre)
    if not nombre:
        return False
    if any(ex in nombre for ex in EXCLUIDOS):
        return False
    return fila.precio_centroabastos is not None or fila.precio_corabastos is not None


def construir_lista_productos(filas: list[FilaSipsa]) -> list[dict]:
    """Lista {x, nombre} de filas que sí son productos, para pasarle al LLM
    normalizador. Se conserva el índice absoluto de fila (x) para poder
    volver a buscarla luego en ``filas``."""
    return [{"x": f.index, "nombre": f.nombre} for f in filas if es_producto(f)]
